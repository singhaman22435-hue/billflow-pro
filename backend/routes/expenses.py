from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models.expense import Expense
from models.vendor import Vendor
from datetime import date, datetime
import io

expenses_bp = Blueprint('expenses', __name__)

def org_id():
    return current_user.org_id

@expenses_bp.route('/expenses')
@login_required
def list():
    category = request.args.get('category', '')
    month = request.args.get('month', date.today().strftime('%Y-%m'))
    try:
        year, mon = map(int, month.split('-'))
        m_start = date(year, mon, 1)
        m_end = date(year, mon + 1, 1) if mon < 12 else date(year + 1, 1, 1)
    except Exception:
        from datetime import timedelta
        m_start = date.today().replace(day=1)
        m_end = date.today() + timedelta(days=1)

    page = request.args.get('page', 1, type=int)

    query = Expense.query.filter(
        Expense.org_id == org_id(),
        Expense.expense_date >= m_start,
        Expense.expense_date < m_end
    )
    if category:
        query = query.filter(Expense.category == category)
    query = query.order_by(Expense.expense_date.desc())
    
    total = sum(float(e.amount or 0) for e in query.all())
    
    pagination = query.paginate(page=page, per_page=50, error_out=False)
    expenses = pagination.items
    
    return render_template('expenses/list.html',
        expenses=expenses, pagination=pagination, total=total,
        categories=Expense.CATEGORIES, selected_cat=category, month=month
    )

@expenses_bp.route('/expenses/add', methods=['GET', 'POST'])
@login_required
def add():
    if request.method == 'POST':
        vendor_id_raw = request.form.get('vendor_id')
        vendor_id = int(vendor_id_raw) if vendor_id_raw and vendor_id_raw.isdigit() else None
        
        e = Expense(
            org_id=org_id(),
            expense_date=datetime.strptime(request.form.get('expense_date'), '%Y-%m-%d').date(),
            category=request.form.get('category'),
            description=request.form.get('description'),
            amount=request.form.get('amount') or 0,
            payment_mode=request.form.get('payment_mode', 'cash'),
            reference_no=request.form.get('reference_no'),
            vendor_id=vendor_id,
            is_gst_applicable=bool(request.form.get('is_gst_applicable')),
            gst_amount=request.form.get('gst_amount') or 0,
            created_by=current_user.id
        )
        db.session.add(e)
        db.session.commit()
        flash('Expense recorded!', 'success')
        return redirect(url_for('expenses.list'))
    vendors = Vendor.query.filter_by(org_id=org_id(), is_active=True).order_by(Vendor.name).all()
    return render_template('expenses/add.html', categories=Expense.CATEGORIES,
                           vendors=vendors, today=date.today())

@expenses_bp.route('/expenses/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit(id):
    e = Expense.query.filter_by(id=id, org_id=org_id()).first_or_404()
    if request.method == 'POST':
        vendor_id_raw = request.form.get('vendor_id')
        e.expense_date = datetime.strptime(request.form.get('expense_date'), '%Y-%m-%d').date()
        e.category = request.form.get('category')
        e.description = request.form.get('description')
        e.amount = request.form.get('amount') or 0
        e.payment_mode = request.form.get('payment_mode', 'cash')
        e.reference_no = request.form.get('reference_no')
        e.vendor_id = int(vendor_id_raw) if vendor_id_raw and vendor_id_raw.isdigit() else None
        e.is_gst_applicable = bool(request.form.get('is_gst_applicable'))
        e.gst_amount = request.form.get('gst_amount') or 0
        db.session.commit()
        flash('Expense updated!', 'success')
        return redirect(url_for('expenses.list'))
    vendors = Vendor.query.filter_by(org_id=org_id(), is_active=True).order_by(Vendor.name).all()
    return render_template('expenses/edit.html', expense=e, categories=Expense.CATEGORIES,
                           vendors=vendors, today=date.today())

@expenses_bp.route('/expenses/<int:id>/delete', methods=['POST'])
@login_required
def delete(id):
    e = Expense.query.filter_by(id=id, org_id=org_id()).first_or_404()
    db.session.delete(e)
    db.session.commit()
    flash('Expense deleted.', 'info')
    return redirect(url_for('expenses.list'))


@expenses_bp.route('/expenses/export')
@login_required
def export_excel():
    """Export expenses as Excel"""
    month = request.args.get('month', date.today().strftime('%Y-%m'))
    try:
        year, mon = map(int, month.split('-'))
        m_start = date(year, mon, 1)
        m_end = date(year, mon + 1, 1) if mon < 12 else date(year + 1, 1, 1)
    except Exception:
        m_start = date.today().replace(day=1)
        m_end = date.today()

    expenses = Expense.query.filter(
        Expense.org_id == org_id(),
        Expense.expense_date >= m_start,
        Expense.expense_date < m_end
    ).order_by(Expense.expense_date).all()

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Expenses'

        # Styles
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill('solid', fgColor='6366F1')
        center = Alignment(horizontal='center')
        right = Alignment(horizontal='right')
        thin = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )

        # Title
        ws.merge_cells('A1:G1')
        ws['A1'] = f'Expense Report — {month}'
        ws['A1'].font = Font(bold=True, size=14, color='1A1A2E')
        ws['A1'].alignment = center

        # Headers
        headers = ['Date', 'Category', 'Description', 'Payment Mode', 'Reference', 'Amount', 'GST']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin

        # Rows
        for row_idx, exp in enumerate(expenses, 4):
            data = [
                exp.expense_date.strftime('%d/%m/%Y'),
                exp.category,
                exp.description or '',
                exp.payment_mode.upper(),
                exp.reference_no or '',
                float(exp.amount),
                float(exp.gst_amount)
            ]
            for col, val in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = thin
                if col >= 6:
                    cell.alignment = right
                    cell.number_format = '#,##0.00'

        # Total row
        total_row = len(expenses) + 4
        ws.cell(row=total_row, column=5, value='TOTAL').font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=sum(float(e.amount or 0) for e in expenses)).font = Font(bold=True)
        ws.cell(row=total_row, column=6).number_format = '#,##0.00'
        ws.cell(row=total_row, column=6).alignment = right

        # Column widths
        col_widths = [12, 20, 30, 14, 20, 14, 10]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=f'expenses_{month}.xlsx')
    except ImportError:
        flash('openpyxl not installed. Run: pip install openpyxl', 'danger')
        return redirect(url_for('expenses.list'))
