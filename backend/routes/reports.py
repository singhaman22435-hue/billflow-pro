from flask import Blueprint, render_template, request, send_file
from flask_login import login_required, current_user
from extensions import db
from models.invoice import Invoice, InvoiceItem
from models.customer import Customer
from models.payment import Payment
from sqlalchemy import func
from datetime import date, datetime
import io

reports_bp = Blueprint('reports', __name__)

def org_id():
    return current_user.org_id

@reports_bp.before_request
def check_reports_permission():
    if not current_user.is_authenticated:
        from flask import redirect, url_for, request
        return redirect(url_for('auth.login', next=request.url))
    if not current_user.can('reports'):
        from flask import flash, redirect, url_for
        flash('You do not have permission to view reports.', 'danger')
        return redirect(url_for('dashboard.index'))

@reports_bp.route('/reports')
@login_required
def index():
    return redirect_to_sales()

def redirect_to_sales():
    from flask import redirect, url_for
    return redirect(url_for('reports.sales'))

@reports_bp.route('/reports/sales')
@login_required
def sales():
    if not current_user.can('reports'):
        from flask import flash, redirect, url_for
        flash('You do not have permission to view reports.', 'danger')
        return redirect(url_for('dashboard.index'))
        
    from_date = request.args.get('from', date.today().replace(day=1).isoformat())
    to_date = request.args.get('to', date.today().isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= fd,
        Invoice.invoice_date <= td,
        Invoice.status != 'cancelled'
    ).order_by(Invoice.invoice_date).all()

    total = sum(float(i.grand_total) for i in invoices)
    total_tax = sum(float(i.cgst_amount) + float(i.sgst_amount) + float(i.igst_amount) for i in invoices)

    return render_template('reports/sales.html', invoices=invoices,
                           from_date=from_date, to_date=to_date,
                           total=total, total_tax=total_tax)

@reports_bp.route('/reports/outstanding')
@login_required
def outstanding():
    """Age-wise receivables analysis"""
    from datetime import timedelta
    today = date.today()

    # Get all unpaid/partial invoices
    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.status.in_(['sent', 'partial']),
        Invoice.balance_due > 0
    ).all()

    # Bucket into age groups
    buckets = {'current': [], 'days_30': [], 'days_60': [], 'days_90': [], 'over_90': []}
    for inv in invoices:
        if not inv.due_date:
            buckets['current'].append(inv)
            continue
        age = (today - inv.due_date).days
        if age <= 0:
            buckets['current'].append(inv)
        elif age <= 30:
            buckets['days_30'].append(inv)
        elif age <= 60:
            buckets['days_60'].append(inv)
        elif age <= 90:
            buckets['days_90'].append(inv)
        else:
            buckets['over_90'].append(inv)

    def total_bucket(b):
        return sum(float(i.balance_due) for i in b)

    return render_template('reports/outstanding.html',
        buckets=buckets, today=today,
        total_outstanding=sum(float(i.balance_due) for i in invoices),
        bucket_totals={k: total_bucket(v) for k, v in buckets.items()}
    )


@reports_bp.route('/reports/purchase')
@login_required
def purchase():
    from_date = request.args.get('from', date.today().replace(day=1).isoformat())
    to_date = request.args.get('to', date.today().isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    from models.purchase_order import PurchaseOrder
    pos = PurchaseOrder.query.filter(
        PurchaseOrder.org_id == org_id(),
        PurchaseOrder.po_date >= fd,
        PurchaseOrder.po_date <= td,
        PurchaseOrder.status != 'cancelled'
    ).order_by(PurchaseOrder.po_date).all()

    total = sum(float(p.grand_total) for p in pos)
    return render_template('reports/purchase.html', pos=pos,
                           from_date=from_date, to_date=to_date, total=total)


@reports_bp.route('/reports/gstr1')
@login_required
def gstr1():
    month = request.args.get('month', date.today().strftime('%Y-%m'))
    year, mon = map(int, month.split('-'))
    from datetime import timedelta
    m_start = date(year, mon, 1)
    if mon == 12:
        m_end = date(year+1, 1, 1)
    else:
        m_end = date(year, mon+1, 1)

    # B2B invoices (customer has GSTIN)
    b2b = db.session.query(Invoice).join(Customer).filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= m_start,
        Invoice.invoice_date < m_end,
        Invoice.status != 'cancelled',
        Customer.gstin != None,
        Customer.gstin != ''
    ).all()

    # B2C invoices
    b2c = db.session.query(Invoice).join(Customer).filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= m_start,
        Invoice.invoice_date < m_end,
        Invoice.status != 'cancelled',
        (Customer.gstin == None) | (Customer.gstin == '')
    ).all()

    return render_template('reports/gstr1.html', b2b=b2b, b2c=b2c, month=month)


# ─── PROFIT & LOSS ─────────────────────────────────────────────────────────────

@reports_bp.route('/reports/pnl')
@login_required
def pnl():
    from models.expense import Expense
    # Default: current financial year
    today = date.today()
    if today.month >= 4:
        default_from = date(today.year, 4, 1).isoformat()
    else:
        default_from = date(today.year - 1, 4, 1).isoformat()
    from_date = request.args.get('from', default_from)
    to_date = request.args.get('to', today.isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= fd, Invoice.invoice_date <= td,
        Invoice.status.in_(['paid', 'partial', 'sent'])
    ).all()

    expenses = Expense.query.filter(
        Expense.org_id == org_id(),
        Expense.expense_date >= fd, Expense.expense_date <= td
    ).all()

    revenue = sum(float(i.taxable_amount) for i in invoices)
    total_gst = sum(float(i.cgst_amount) + float(i.sgst_amount) + float(i.igst_amount) for i in invoices)
    total_revenue_incl_tax = sum(float(i.grand_total) for i in invoices)
    total_expenses = sum(float(e.amount or 0) for e in expenses)

    expense_by_cat = {}
    for e in expenses:
        expense_by_cat[e.category] = expense_by_cat.get(e.category, 0) + float(e.amount or 0)

    net_profit = revenue - total_expenses

    # Month-wise data for chart
    monthly_keys = []
    monthly_revenue = []
    monthly_expense = []
    monthly_map = {}
    for inv in invoices:
        k = inv.invoice_date.strftime('%b %Y')
        if k not in monthly_map:
            monthly_map[k] = {'rev': 0, 'exp': 0}
        monthly_map[k]['rev'] += float(inv.taxable_amount)
    for e in expenses:
        k = e.expense_date.strftime('%b %Y')
        if k not in monthly_map:
            monthly_map[k] = {'rev': 0, 'exp': 0}
        monthly_map[k]['exp'] += float(e.amount or 0)
    for k in sorted(monthly_map.keys(), key=lambda x: datetime.strptime(x, '%b %Y')):
        monthly_keys.append(k)
        monthly_revenue.append(round(monthly_map[k]['rev'], 2))
        monthly_expense.append(round(monthly_map[k]['exp'], 2))

    return render_template('reports/pnl.html',
        revenue=revenue, total_gst=total_gst,
        total_revenue_incl_tax=total_revenue_incl_tax,
        total_expenses=total_expenses, expense_by_cat=expense_by_cat,
        net_profit=net_profit, from_date=from_date, to_date=to_date,
        invoice_count=len(invoices), expense_count=len(expenses),
        monthly_keys=monthly_keys, monthly_revenue=monthly_revenue,
        monthly_expense=monthly_expense
    )


# ─── CASH BOOK ─────────────────────────────────────────────────────────────────

@reports_bp.route('/reports/cashbook')
@login_required
def cashbook():
    from models.expense import Expense
    month = request.args.get('month', date.today().strftime('%Y-%m'))
    try:
        year, mon = map(int, month.split('-'))
        m_start = date(year, mon, 1)
        m_end = date(year, mon + 1, 1) if mon < 12 else date(year + 1, 1, 1)
    except Exception:
        m_start = date.today().replace(day=1)
        m_end = date.today()

    payments_in = Payment.query.join(Invoice).filter(
        Invoice.org_id == org_id(),
        Payment.payment_date >= m_start,
        Payment.payment_date < m_end
    ).order_by(Payment.payment_date).all()

    payments_out = Expense.query.filter(
        Expense.org_id == org_id(),
        Expense.expense_date >= m_start,
        Expense.expense_date < m_end
    ).order_by(Expense.expense_date).all()

    ledger = []
    for p in payments_in:
        ledger.append({
            'date': p.payment_date,
            'type': 'in',
            'narration': f"Received — {p.invoice.invoice_no} ({p.invoice.customer.name})",
            'mode': p.payment_mode.upper(),
            'ref': p.reference_no or '',
            'debit': 0.0,
            'credit': float(p.amount)
        })
    for e in payments_out:
        ledger.append({
            'date': e.expense_date,
            'type': 'out',
            'narration': f"{e.category} — {e.description or ''}",
            'mode': e.payment_mode.upper(),
            'ref': e.reference_no or '',
            'debit': float(e.amount or 0),
            'credit': 0.0
        })

    ledger.sort(key=lambda x: x['date'])
    balance = 0
    for entry in ledger:
        balance += entry['credit'] - entry['debit']
        entry['balance'] = round(balance, 2)

    total_in = sum(e['credit'] for e in ledger)
    total_out = sum(e['debit'] for e in ledger)

    return render_template('reports/cashbook.html',
        ledger=ledger, month=month,
        total_in=total_in, total_out=total_out,
        closing_balance=total_in - total_out
    )


# ─── EXCEL EXPORTS ─────────────────────────────────────────────────────────────

@reports_bp.route('/reports/export/invoices')
@login_required
def export_invoices():
    from_date = request.args.get('from', date.today().replace(day=1).isoformat())
    to_date = request.args.get('to', date.today().isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= fd, Invoice.invoice_date <= td
    ).order_by(Invoice.invoice_date).all()

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Invoices'
        h_font = Font(bold=True, color='FFFFFF', size=10)
        h_fill = PatternFill('solid', fgColor='6366F1')
        right = Alignment(horizontal='right')
        center = Alignment(horizontal='center')
        thin = Border(left=Side(style='thin', color='E5E7EB'), right=Side(style='thin', color='E5E7EB'),
                      top=Side(style='thin', color='E5E7EB'), bottom=Side(style='thin', color='E5E7EB'))

        ws.merge_cells('A1:L1')
        ws['A1'] = f'Invoice Register: {from_date} to {to_date}'
        ws['A1'].font = Font(bold=True, size=13)

        headers = ['#', 'Invoice No', 'Date', 'Customer', 'GSTIN', 'State',
                   'Taxable', 'CGST', 'SGST', 'IGST', 'Grand Total', 'Status']
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col, value=h)
            c.font = h_font; c.fill = h_fill; c.alignment = center; c.border = thin

        for i, inv in enumerate(invoices, 1):
            row_data = [i, inv.invoice_no, inv.invoice_date.strftime('%d/%m/%Y'),
                        inv.customer.name, inv.customer.gstin or '', inv.customer.state or '',
                        float(inv.taxable_amount), float(inv.cgst_amount), float(inv.sgst_amount),
                        float(inv.igst_amount), float(inv.grand_total), inv.status]
            for col, val in enumerate(row_data, 1):
                c = ws.cell(row=i + 3, column=col, value=val)
                c.border = thin
                if 7 <= col <= 11:
                    c.alignment = right; c.number_format = '#,##0.00'

        tr = len(invoices) + 4
        ws.cell(row=tr, column=6, value='TOTAL').font = Font(bold=True)
        for col, attr in [(7,'taxable_amount'),(8,'cgst_amount'),(9,'sgst_amount'),(10,'igst_amount'),(11,'grand_total')]:
            c = ws.cell(row=tr, column=col, value=sum(float(getattr(i, attr)) for i in invoices))
            c.font = Font(bold=True); c.number_format = '#,##0.00'; c.alignment = right

        for i, w in enumerate([4,14,12,24,18,14,12,10,10,10,12,10], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=f'invoices_{from_date}_{to_date}.xlsx')
    except ImportError:
        from flask import flash, redirect, url_for
        flash('openpyxl not installed. Run: pip install openpyxl', 'danger')
        return redirect(url_for('reports.sales'))


@reports_bp.route('/reports/export/pnl')
@login_required
def export_pnl():
    from models.expense import Expense
    today = date.today()
    from_date = request.args.get('from', date(today.year if today.month >= 4 else today.year-1, 4, 1).isoformat())
    to_date = request.args.get('to', today.isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= fd, Invoice.invoice_date <= td,
        Invoice.status.in_(['paid', 'partial', 'sent'])
    ).all()
    expenses = Expense.query.filter(
        Expense.org_id == org_id(),
        Expense.expense_date >= fd, Expense.expense_date <= td
    ).all()

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'P&L Statement'
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 18

        def add_row(label, value=None, bold=False, fill_hex=None):
            r = ws.max_row + 1
            ca = ws.cell(row=r, column=1, value=label)
            ca.font = Font(bold=bold, size=10)
            if fill_hex:
                ca.fill = PatternFill('solid', fgColor=fill_hex)
            if value is not None:
                cb = ws.cell(row=r, column=2, value=value)
                cb.font = Font(bold=bold, size=10)
                cb.alignment = Alignment(horizontal='right')
                cb.number_format = '#,##0.00'
                if fill_hex:
                    cb.fill = PatternFill('solid', fgColor=fill_hex)

        ws.cell(row=1, column=1, value='Profit & Loss Statement').font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value=f'{from_date} to {to_date}')
        ws.append([])

        revenue = sum(float(i.taxable_amount) for i in invoices)
        add_row('INCOME', bold=True, fill_hex='E0E7FF')
        add_row('Sales Revenue (Taxable)', revenue)
        add_row('GST Collected', sum(float(i.cgst_amount)+float(i.sgst_amount)+float(i.igst_amount) for i in invoices))
        add_row('Total Invoiced', sum(float(i.grand_total) for i in invoices), bold=True)
        ws.append([])

        exp_by_cat = {}
        for e in expenses:
            exp_by_cat[e.category] = exp_by_cat.get(e.category, 0) + float(e.amount or 0)
        total_exp = sum(exp_by_cat.values())

        add_row('EXPENSES', bold=True, fill_hex='FEE2E2')
        for cat, amt in sorted(exp_by_cat.items()):
            add_row(f'  {cat}', amt)
        add_row('Total Expenses', total_exp, bold=True)
        ws.append([])

        net = revenue - total_exp
        add_row('NET PROFIT / LOSS', net, bold=True, fill_hex='D1FAE5' if net >= 0 else 'FEE2E2')

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=f'pnl_{from_date}_{to_date}.xlsx')
    except ImportError:
        from flask import flash, redirect, url_for
        flash('openpyxl not installed.', 'danger')
        return redirect(url_for('reports.pnl'))


# ─── GSTR-3B SUMMARY ────────────────────────────────────────
@reports_bp.route('/reports/gstr3b')
@login_required
def gstr3b():
    from flask import redirect, url_for
    from_date = request.args.get('from', date.today().replace(day=1).isoformat())
    to_date = request.args.get('to', date.today().isoformat())
    fd = datetime.strptime(from_date, '%Y-%m-%d').date()
    td = datetime.strptime(to_date, '%Y-%m-%d').date()

    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.invoice_date >= fd,
        Invoice.invoice_date <= td,
        Invoice.status.in_(['sent', 'partial', 'paid']),
        Invoice.invoice_type == 'tax_invoice'
    ).all()

    # 3.1 — Outward taxable supplies
    total_taxable = sum(float(i.taxable_amount) for i in invoices)
    total_cgst = sum(float(i.cgst_amount) for i in invoices)
    total_sgst = sum(float(i.sgst_amount) for i in invoices)
    total_igst = sum(float(i.igst_amount) for i in invoices)
    total_cess = sum(float(i.cess_amount) for i in invoices)
    total_tax = total_cgst + total_sgst + total_igst + total_cess
    total_turnover = sum(float(i.grand_total) for i in invoices)

    # Breakup by GST rate
    from models.invoice import InvoiceItem
    rate_wise = {}
    for inv in invoices:
        for item in inv.items:
            rate = float(item.gst_rate or 0)
            if rate not in rate_wise:
                rate_wise[rate] = {'taxable': 0, 'cgst': 0, 'sgst': 0, 'igst': 0}
            rate_wise[rate]['taxable'] += float(item.taxable_amount)
            if inv.is_igst:
                rate_wise[rate]['igst'] += float(item.igst_amount)
            else:
                rate_wise[rate]['cgst'] += float(item.cgst_amount)
                rate_wise[rate]['sgst'] += float(item.sgst_amount)

    return render_template('reports/gstr3b.html',
        from_date=from_date, to_date=to_date,
        invoices=invoices,
        total_taxable=total_taxable,
        total_cgst=total_cgst, total_sgst=total_sgst,
        total_igst=total_igst, total_cess=total_cess,
        total_tax=total_tax, total_turnover=total_turnover,
        rate_wise=dict(sorted(rate_wise.items())))


# ─── INVOICE AGING REPORT ───────────────────────────────────
@reports_bp.route('/reports/aging')
@login_required
def aging():
    today = date.today()
    invoices = Invoice.query.filter(
        Invoice.org_id == org_id(),
        Invoice.status.in_(['sent', 'partial']),
        Invoice.balance_due > 0
    ).all()

    buckets = {
        'current': [],    # not yet due
        '1_30': [],       # 1-30 days overdue
        '31_60': [],      # 31-60 days overdue
        '61_90': [],      # 61-90 days overdue
        '90_plus': []     # 90+ days overdue
    }

    for inv in invoices:
        if not inv.due_date:
            buckets['current'].append(inv)
            continue
        days = (today - inv.due_date).days
        if days <= 0:
            buckets['current'].append(inv)
        elif days <= 30:
            buckets['1_30'].append(inv)
        elif days <= 60:
            buckets['31_60'].append(inv)
        elif days <= 90:
            buckets['61_90'].append(inv)
        else:
            buckets['90_plus'].append(inv)

    totals = {k: sum(float(i.balance_due) for i in v) for k, v in buckets.items()}
    grand_total = sum(totals.values())

    return render_template('reports/aging.html',
        buckets=buckets, totals=totals, grand_total=grand_total, today=today)

